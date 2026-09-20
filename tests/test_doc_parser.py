"""文件解析器测试：PDF / Word / Excel / 文本，均在内存中构造样本。"""
from __future__ import annotations

import io
import sys
import unittest
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from src.services import doc_parser  # noqa: E402


def _make_docx() -> bytes:
    import docx

    document = docx.Document()
    document.add_heading("第三章 虚拟语气", level=1)
    document.add_paragraph("虚拟语气用于表达与事实相反的假设。")
    document.add_heading("3.2 用法", level=2)
    document.add_paragraph("主句用 would + 动词原形。")
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _make_xlsx() -> bytes:
    import openpyxl

    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "题库"
    sheet.append(["题干", "选项", "正确答案", "解析"])
    sheet.append(["虚拟语气表示什么？", "A 事实 B 假设", "B", "表达与事实相反"])
    sheet.append(["wish 后面用什么时态？", "A 过去式 B 将来", "A", "虚拟语气"])
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _make_pdf() -> bytes:
    import fitz

    document = fitz.open()
    page = document.new_page()
    # 必须用内置中文字体，默认字体渲染中文会变成省略号
    page.insert_text((72, 100), "第三章 虚拟语气", fontsize=16, fontname="china-s")
    page.insert_text((72, 130), "虚拟语气用于表达与事实相反的假设。", fontsize=11, fontname="china-s")
    data = document.tobytes()
    document.close()
    return data


class TestDocParser(unittest.TestCase):
    def test_text_passthrough(self) -> None:
        self.assertEqual(doc_parser.parse_file("a.md", "# 标题".encode()), "# 标题")

    def test_text_tolerates_non_utf8(self) -> None:
        result = doc_parser.parse_file("a.txt", b"\xff\xfe hello")
        self.assertIn("hello", result)

    def test_docx_heading_levels(self) -> None:
        markdown = doc_parser.parse_file("讲义.docx", _make_docx())
        self.assertIn("# 第三章 虚拟语气", markdown)
        self.assertIn("## 3.2 用法", markdown)
        self.assertIn("虚拟语气用于表达与事实相反的假设。", markdown)

    def test_xlsx_becomes_structured_questions(self) -> None:
        markdown = doc_parser.parse_file("题库.xlsx", _make_xlsx())
        self.assertIn("# 题库", markdown)
        self.assertIn("## 题目 1", markdown)
        self.assertIn("**题干**：虚拟语气表示什么？", markdown)
        self.assertIn("**正确答案**：B", markdown)
        self.assertIn("**解析**：表达与事实相反", markdown)

    def test_xlsx_without_question_header_fails_loudly(self) -> None:
        import openpyxl

        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.append(["A", "B", "C"])
        sheet.append([1, 2, 3])
        buffer = io.BytesIO()
        workbook.save(buffer)
        with self.assertRaises(ValueError) as ctx:
            doc_parser.parse_file("乱表.xlsx", buffer.getvalue())
        self.assertIn("题干", str(ctx.exception))

    def test_pdf_text_extraction(self) -> None:
        markdown = doc_parser.parse_file("手册.pdf", _make_pdf())
        self.assertIn("虚拟语气", markdown)
        self.assertIn("第 1 页", markdown)

    def test_unsupported_suffix(self) -> None:
        with self.assertRaises(doc_parser.UnsupportedFormatError):
            doc_parser.parse_file("图片.png", b"123")

    def test_empty_pdf_fails(self) -> None:
        import fitz

        document = fitz.open()
        document.new_page()
        data = document.tobytes()
        document.close()
        with self.assertRaises(ValueError):
            doc_parser.parse_file("空白.pdf", data)


if __name__ == "__main__":
    unittest.main()
